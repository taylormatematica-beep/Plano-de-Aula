"""
Assistente de Plano de Aula — Escola Presidente Bernardes
Execute:  python app.py   e acesse http://localhost:5000
"""
import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from io import BytesIO

from gerador import CONFIG_FILE, carregar_config, gerar_plano
from pdf import gerar_pdf
from pdf_atividade import gerar_pdf_atividade, gerar_pdf_gabarito
from gerador import gerar_atividade, TIPOS_ATIVIDADE, distribuir_valores
import correcao
import db
import drive
import fuso
import referencias

db.init()

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False

# ----------------------------------------------------------------------------
# Autenticação por usuário (e-mail institucional + senha)
# ----------------------------------------------------------------------------
import os, hmac  # noqa: E402
from flask import redirect, session, url_for, abort, g  # noqa: E402
import auth  # noqa: E402
import email_util  # noqa: E402


def _env(nome: str) -> str:
    """Lê variável de ambiente removendo espaços e aspas acidentais."""
    return os.getenv(nome, "").strip().strip('"').strip("'").strip()


DIRECAO_NOME = _env("DIRECAO_NOME")           # nome impresso no campo Direção
SUPERVISAO_NOME = _env("SUPERVISAO_NOME")     # nome padrão da supervisão
SUPERVISAO_SENHA = _env("SUPERVISAO_SENHA")   # senha de emergência da supervisão (opcional)
PUBLIC_URL = _env("PUBLIC_URL")

app.secret_key = os.getenv("SECRET_KEY") or (SUPERVISAO_SENHA + "-plano-bernardes-2026")
app.config.update(
    SESSION_COOKIE_NAME="planoaula_sessao",
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SECURE=bool(os.getenv("RENDER") or os.getenv("FORCE_HTTPS")),  # https no Render
    PERMANENT_SESSION_LIFETIME=60 * 60 * 24 * 14,  # 14 dias
)

from werkzeug.middleware.proxy_fix import ProxyFix  # noqa: E402
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

ROTAS_PUBLICAS = {"login", "primeiro_acesso", "definir_senha", "static", "sup_login", "healthz",
                  "prova_entrar", "prova_aluno", "api_prova_aluno_info", "api_prova_aluno_enviar"}


def _base_url() -> str:
    base = PUBLIC_URL or request.url_root.rstrip("/")
    if base.startswith("http://") and "localhost" not in base and "127.0.0.1" not in base:
        base = "https://" + base[len("http://"):]
    return base


def usuario_atual() -> dict | None:
    """Usuário logado. Consulta o banco UMA vez por requisição (cache em flask.g)."""
    uid = session.get("uid")
    if not uid:
        return None
    if "usuario_cache" in g and g.usuario_cache_id == uid:
        return g.usuario_cache
    u = db.usuario_por_id(uid)
    if not u or not u.get("ativo", 1):
        session.clear()
        return None
    g.usuario_cache, g.usuario_cache_id = u, uid
    return u


def _e_supervisao() -> bool:
    if session.get("sup"):
        return True
    u = usuario_atual()
    return bool(u and u.get("perfil") == "supervisao")


def _nome_professor() -> str:
    u = usuario_atual()
    return (u or {}).get("nome") or session.get("professor", "")


def _email_professor() -> str:
    u = usuario_atual()
    return (u or {}).get("email", "")


@app.before_request
def _exigir_login():
    if request.endpoint in ROTAS_PUBLICAS or (request.endpoint or "").startswith("static"):
        return None
    if usuario_atual() or session.get("sup"):
        return None
    if request.path.startswith("/api/"):
        return jsonify({"erro": "Sessão expirada. Entre novamente.", "login": True}), 401
    return redirect(url_for("login", next=request.path))


@app.route("/healthz")
def healthz():
    """Verificação pública (sem segredos): arquivos, banco e configuração básica."""
    import importlib, traceback as tb
    linhas = []
    base = Path(__file__).parent
    for f in ["auth.py", "email_util.py", "db.py", "drive.py", "gerador.py", "pdf.py",
              "templates/login.html", "templates/index.html", "templates/historico.html",
              "templates/conta.html", "templates/usuarios.html", "static/logo.png"]:
        linhas.append(f"{'OK ' if (base / f).exists() else 'FALTA'}  {f}")
    try:
        with db.conexao() as con:
            con.execute("SELECT 1")
        linhas.append(f"OK   banco de dados ({'PostgreSQL' if db.USA_PG else 'SQLite local'})")
        try:
            n = len(db.usuarios_listar())
            linhas.append(f"OK   tabela usuarios ({n} usuário(s))")
        except Exception as e:  # noqa: BLE001
            linhas.append(f"ERRO tabela usuarios: {type(e).__name__}: {str(e)[:160]}")
    except Exception as e:  # noqa: BLE001
        linhas.append(f"ERRO banco de dados: {type(e).__name__}: {str(e)[:160]}")
    linhas.append(f"{'OK ' if email_util.configurado() else 'AVISO'}  E-mail: {email_util.descricao() if email_util.configurado() else 'não configurado (links aparecerão na tela da supervisão)'}")
    linhas.append(f"{'OK ' if auth.SUPERVISAO_EMAILS else 'AVISO'}  SUPERVISAO_EMAILS {'definido' if auth.SUPERVISAO_EMAILS else 'vazio'}")
    linhas.append(f"{'OK ' if PUBLIC_URL else 'AVISO'}  PUBLIC_URL {'definido' if PUBLIC_URL else 'vazio (usará o endereço da requisição)'}")
    try:
        render_template("login.html", modo="login", erro="", dominio=auth.dominio_msg(), email="", ok="")
        linhas.append("OK   template login.html renderiza")
    except Exception as e:  # noqa: BLE001
        linhas.append(f"ERRO ao renderizar login.html: {type(e).__name__}: {str(e)[:160]}")
    status = 200 if not any(l.startswith("ERRO") or l.startswith("FALTA") for l in linhas) else 500
    return "<pre style='font:14px/1.6 monospace;padding:16px'>VERIFICAÇÃO DO SISTEMA\n\n" + "\n".join(linhas) + "</pre>", status


@app.errorhandler(500)
def _erro_500(e):
    import traceback
    tb = traceback.format_exc()
    app.logger.error("ERRO 500 em %s\n%s", request.path, tb)
    print("ERRO 500 em", request.path, "\n", tb, flush=True)
    ultima = [l for l in tb.strip().splitlines() if l.strip()][-1] if tb.strip() and tb.strip() != "NoneType: None" else str(e)
    return (f"<!doctype html><meta charset=utf-8><body style='font-family:Segoe UI,Arial;padding:28px;max-width:720px'>"
            f"<h2>Ocorreu um erro no servidor</h2>"
            f"<p style='background:#fdecec;border:1px solid #f3b4b4;padding:10px;border-radius:8px;font-family:monospace;font-size:13px'>{ultima}</p>"
            f"<p>Abra <a href='/healthz'>/healthz</a> para ver a verificação do sistema. O detalhe completo está no log do servidor.</p>"
            f"<p><a href='/login'>← Voltar</a></p>"), 500


@app.route("/login", methods=["GET", "POST"])
def login():
    erro = ""
    if request.method == "POST":
        try:
            u = auth.autenticar(request.form.get("email", ""), request.form.get("senha", ""))
            session.clear()
            session.permanent = bool(request.form.get("lembrar"))
            session["uid"] = u["id"]
            session["professor"] = u.get("nome") or ""
            destino = request.args.get("next") or "/"
            return redirect(destino if destino.startswith("/") else "/")
        except auth.AuthErro as e:
            erro = str(e)
    return render_template("login.html", modo="login", erro=erro, dominio=auth.dominio_msg(),
                           email=request.form.get("email", ""), ok="")


@app.route("/primeiro-acesso", methods=["GET", "POST"])
def primeiro_acesso():
    erro, ok, link_manual = "", "", ""
    if request.method == "POST":
        try:
            r = auth.solicitar_acesso(request.form.get("email", ""), _base_url())
            if r["enviado"]:
                ok = ("Enviamos um link para o seu e-mail. Abra a mensagem e clique em "
                      "<b>Criar minha senha</b>. Verifique também a pasta de spam.")
            else:
                erro = ("Não foi possível enviar o e-mail agora. Procure a supervisão para receber o link de acesso."
                        + (f" (Detalhe técnico: {r['erro']})" if _e_supervisao() else ""))
                if _e_supervisao():
                    link_manual = r["link"] or ""
        except auth.AuthErro as e:
            erro = str(e)
    return render_template("login.html", modo="primeiro", erro=erro, ok=ok, dominio=auth.dominio_msg(),
                           email=request.form.get("email", ""), link_manual=link_manual)


@app.route("/senha/<token>", methods=["GET", "POST"])
def definir_senha(token):
    erro, usuario = "", None
    try:
        t = auth.validar_token(token)
        usuario = db.usuario_por_id(t["usuario_id"])
    except auth.AuthErro as e:
        return render_template("login.html", modo="token_invalido", erro=str(e), dominio=auth.dominio_msg(), ok="")
    if request.method == "POST":
        s1, s2 = request.form.get("senha", ""), request.form.get("senha2", "")
        nome = request.form.get("nome", "").strip()
        if s1 != s2:
            erro = "As senhas não coincidem."
        elif not nome:
            erro = "Informe seu nome completo (ele aparece no plano de aula)."
        else:
            try:
                u = auth.definir_senha(token, s1, nome)
                session.clear()
                session.permanent = True
                session["uid"] = u["id"]
                session["professor"] = u.get("nome") or ""
                return redirect("/?bemvindo=1")
            except auth.AuthErro as e:
                erro = str(e)
    return render_template("login.html", modo="senha", erro=erro, ok="", usuario=usuario, token=token,
                           dominio=auth.dominio_msg(), nome=request.form.get("nome") or usuario.get("nome") or "")


@app.route("/supervisao/login", methods=["GET", "POST"])
def sup_login():
    """Acesso de emergência da supervisão por senha única (SUPERVISAO_SENHA). Opcional."""
    if not SUPERVISAO_SENHA:
        return redirect(url_for("login"))
    erro = ""
    if request.method == "POST":
        digitada = request.form.get("senha", "").strip().encode()
        if hmac.compare_digest(digitada, SUPERVISAO_SENHA.encode()):
            session.permanent = True
            session["sup"] = True
            session.setdefault("professor", SUPERVISAO_NOME or "Supervisão")
            return redirect("/historico")
        erro = "Senha incorreta."
    return render_template("login.html", modo="sup", erro=erro, ok="", dominio=auth.dominio_msg())


@app.route("/supervisao/sair")
def sup_sair():
    session.pop("sup", None)
    return redirect("/historico")


@app.route("/diagnostico")
def diagnostico():
    if not _e_supervisao():
        return "Apenas a supervisão pode ver o diagnóstico.", 403

    def info_senha(v: str) -> str:
        if not v:
            return "❌ NÃO definida"
        return f"✅ definida — {len(v)} caracteres, começa com «{v[0]}» e termina com «{v[-1]}»"
    cfg = carregar_config()
    linhas = {
        "Login por usuário": f"domínio {auth.dominio_msg()} · {len(db.usuarios_listar())} usuário(s) · {db.usuarios_total_supervisao()} supervisão",
        "E-mail": (f"✅ {email_util.descricao()}" if email_util.configurado()
                   else "❌ não configurado (BREVO_API_KEY + EMAIL_FROM) — links terão de ser repassados manualmente"),
        "SUPERVISAO_EMAILS": ", ".join(sorted(auth.SUPERVISAO_EMAILS)) or "— (nenhum; use SUPERVISAO_SENHA ou promova pela tela Usuários)",
        "SUPERVISAO_SENHA (emergência)": info_senha(SUPERVISAO_SENHA),
        "PUBLIC_URL": PUBLIC_URL or "⚠️ vazio (usará o endereço da requisição)",
        "SUPERVISAO_NOME": SUPERVISAO_NOME or "❌ vazio",
        "DIRECAO_NOME": DIRECAO_NOME or "❌ vazio",
        "DATABASE_URL": "✅ PostgreSQL" if db.USA_PG else "⚠️ não definida (usando SQLite local — histórico some a cada deploy no Render)",
        "IA": f"{cfg['provider']} · {cfg['model']} · chave {'✅' if cfg['api_key'] else '❌'}",
        "Google Drive": ("✅ conectado como " + drive.status()["conta"]) if drive.conectado()
                        else ("⚠️ credenciais OK, falta conectar (Configurações)" if drive.credenciais_ok()
                              else "— não configurado (opcional)"),
        "Horário do sistema": f"{fuso.agora().strftime('%d/%m/%Y %H:%M')} ({fuso.FUSO.key}) · UTC do servidor: {__import__('datetime').datetime.now(__import__('datetime').timezone.utc).strftime('%H:%M')}",
        "Sessão atual": f"usuário: {_email_professor() or '—'} · supervisão: {'sim' if _e_supervisao() else 'não'}",
        "Tempo de conexão ao banco": _medir_banco(),
        "Servidor": f"{os.getenv('SERVER_SOFTWARE', '') or 'gunicorn'} · workers/threads conforme startCommand",
    }
    html = "".join(f"<tr><td style='padding:6px 12px;font-weight:600'>{k}</td><td style='padding:6px 12px'>{v}</td></tr>" for k, v in linhas.items())
    return (f"<!doctype html><meta charset=utf-8><title>Diagnóstico</title>"
            f"<body style='font-family:Segoe UI,Arial;padding:24px'><h2>Diagnóstico do servidor</h2>"
            f"<table style='border-collapse:collapse;background:#f6f6f8;border-radius:8px'>{html}</table>"
            f"<p><a href='/healthz'>Verificação de arquivos e banco</a> · <a href='/'>← voltar</a></p>")


@app.route("/sair")
def sair():
    session.clear()
    return redirect(url_for("login"))


@app.route("/conta", methods=["GET", "POST"])
def conta():
    u = usuario_atual()
    if not u:
        return redirect(url_for("login"))
    erro = ok = ""
    if request.method == "POST":
        try:
            if request.form.get("acao") == "nome":
                nome = request.form.get("nome", "").strip()
                if len(nome) < 3:
                    raise auth.AuthErro("Informe seu nome completo.")
                db.usuario_atualizar(u["id"], nome=nome)
                session["professor"] = nome
                ok = "Nome atualizado."
            else:
                auth.trocar_senha(u["id"], request.form.get("atual", ""), request.form.get("nova", ""))
                ok = "Senha alterada com sucesso."
        except auth.AuthErro as e:
            erro = str(e)
        u = usuario_atual()
    return render_template("conta.html", u=u, erro=erro, ok=ok, e_supervisao=_e_supervisao())


# Componentes curriculares conforme as matrizes 2026 da escola (Integral Profissional)
DISCIPLINAS_GRUPOS = {
    "Formação Geral Básica": [
        "Língua Portuguesa", "Língua Inglesa", "Arte", "Educação Física", "Matemática",
        "Física", "Química", "Biologia", "História", "Geografia", "Filosofia", "Sociologia",
    ],
    "Parte Diversificada / Escola da Escolha": [
        "Projeto de Vida", "Eletiva", "Estudos Orientados", "Práticas Experimentais",
        "Nivelamento – Língua Portuguesa", "Nivelamento – Matemática",
        "Cultura Digital e Fundamentos de IA", "Ferramentas para o Mundo do Trabalho",
        "Projetos Integradores e de Corresponsabilidade Social (PICS)",
        "Práticas de Leitura e Escrita",
    ],
    "Técnico – Automação Industrial (2º ano)": ["Automação Industrial III", "Automação Industrial IV"],
    "Técnico – Mecatrônica (2º ano)": ["Mecatrônica III", "Mecatrônica IV"],
    "Técnico – Desenvolvimento de Sistemas (3º ano)": [
        "Conceitos Avançados em Arquitetura de Sistemas", "Desenvolvimento Back-end",
        "Desenvolvimento de Aplicativos", "Desenvolvimento de Softwares", "Desenvolvimento Front-end II",
        "Fundamentos de Segurança de Softwares", "Prática Profissional e Empreendedora",
    ],
}
DISCIPLINAS = [d for grupo in DISCIPLINAS_GRUPOS.values() for d in grupo]
SERIES = ["1º ano", "2º ano", "3º ano"]


def _slug(t: str) -> str:
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-zA-Z0-9]+", "-", t).strip("-").lower()


def _formatar_data(iso: str) -> str:
    """'2026-09-14' -> '14/09/2026'. Aceita também texto livre (ex.: '14 a 18/09/2026')."""
    try:
        y, m, d = iso.split("-")
        return f"{d}/{m}/{y}"
    except Exception:
        return iso


@app.route("/")
def index():
    cfg = carregar_config()
    return render_template(
        "index.html",
        disciplinas=DISCIPLINAS,
        disciplinas_grupos=DISCIPLINAS_GRUPOS,
        series=SERIES,
        hoje=fuso.hoje().isoformat(),
        ia_configurada=bool(cfg["api_key"]),
        provider=cfg["provider"],
        model=cfg["model"],
        supervisao_nome=SUPERVISAO_NOME,
        direcao_nome=DIRECAO_NOME,
        e_supervisao=_e_supervisao(),
        drive_conectado=drive.conectado(),
        usuario=usuario_atual() or {"nome": session.get("professor", ""), "email": ""},
        fontes=referencias.FONTES,
        documentos=db.documentos_listar(somente_ativos=True),
    )


@app.route("/api/verificar-duplicado", methods=["POST"])
def api_verificar_duplicado():
    """Antes de gerar: há outro plano deste professor com o mesmo contexto?"""
    d = request.get_json(force=True) or {}
    u = usuario_atual()
    prof = (u.get("nome") if u else d.get("professor", "")) or ""
    email = (u.get("email") if u else "") or None
    data_ref = _formatar_data(d.get("data", ""))
    if d.get("data_fim"):
        data_ref = f"{data_ref} a {_formatar_data(d['data_fim'])}"
    iguais = db.semelhantes(prof, email, d.get("disciplina", ""), d.get("serie", ""), data_ref, d.get("conteudo", ""))
    return jsonify([{k: i.get(k) for k in ("id", "criado_em", "tema", "data_ref", "conteudo", "visto_em", "drive_link", "motivo")}
                    for i in iguais])


@app.route("/api/gerar", methods=["POST"])
def api_gerar():
    dados = request.get_json(force=True) or {}
    u = usuario_atual()
    if u:
        dados["professor"] = u.get("nome") or dados.get("professor", "")
        dados["professor_email"] = u.get("email", "")
    obrig = ["professor", "disciplina", "conteudo", "data", "serie"]
    faltando = [c for c in obrig if not str(dados.get(c, "")).strip()]
    if faltando:
        return jsonify({"erro": f"Preencha: {', '.join(faltando)}"}), 400

    dados["data"] = _formatar_data(dados["data"])
    if dados.get("data_fim"):
        dados["data"] = f"{dados['data']} a {_formatar_data(dados['data_fim'])}"
    fontes_sel = dados.get("fontes") or ["bncc", "crmg", "ice"]
    if isinstance(fontes_sel, str):
        fontes_sel = [f for f in fontes_sel.split(",") if f]
    doc_ids = [int(x) for x in (dados.get("documentos") or []) if str(x).isdigit()]
    consulta = f"{dados.get('conteudo','')} {dados.get('disciplina','')}"
    try:
        contexto, citacoes = referencias.montar_contexto(fontes_sel, doc_ids, consulta)
    except Exception:  # noqa: BLE001
        contexto, citacoes = "", []
    try:
        plano = gerar_plano(dados, contexto=contexto, citacoes=citacoes)
    except Exception as e:  # noqa: BLE001
        msg = re.sub(r"[?&]key=[^&\s)\]]+", "?key=***", str(e))  # nunca expor a chave
        return jsonify({"erro": f"Falha ao consultar a IA: {msg}"}), 502
    dados.setdefault("direcao", DIRECAO_NOME)
    if not dados.get("supervisao"):
        dados["supervisao"] = SUPERVISAO_NOME
    session["professor"] = dados["professor"].strip()
    id_ = db.inserir(dados, plano)
    return jsonify({"id": id_, "dados": dados, "plano": plano})


@app.route("/api/historico/<int:id_>/salvar", methods=["POST"])
def api_historico_salvar(id_):
    """Salva as edições feitas na tela. Se o plano tinha visto, o visto é removido (o conteúdo mudou)."""
    item = db.obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver(item):
        return jsonify({"erro": "sem permissão"}), 403
    body = request.get_json(force=True) or {}
    dados, plano = body.get("dados") or item["dados"], body.get("plano") or item["plano"]
    plano = {k: v for k, v in plano.items() if not k.startswith("_")}
    # mantém campos internos (e-mail do professor etc.)
    for k in ("professor_email",):
        if item["dados"].get(k) and not dados.get(k):
            dados[k] = item["dados"][k]
    tinha_visto = bool(item.get("visto_em"))
    db.atualizar(id_, dados, plano)
    if tinha_visto:
        db.marcar_visto(id_, "", desfazer=True)
    # se já está no Drive, atualiza o arquivo (sem carimbo agora)
    aviso = ""
    if item.get("drive_file_id") and drive.conectado():
        try:
            _enviar_drive(id_, dados, plano)
        except Exception as e:  # noqa: BLE001
            aviso = f"Salvo, mas não foi possível atualizar o PDF no Drive: {e}"
    return jsonify({"ok": True, "visto_removido": tinha_visto, "aviso": aviso})


@app.route("/api/pdf", methods=["POST"])
def api_pdf():
    body = request.get_json(force=True) or {}
    dados, plano = body.get("dados", {}), body.get("plano", {})
    plano = {k: v for k, v in plano.items() if not k.startswith("_")}
    visto = None
    if body.get("id"):
        item = db.obter(int(body["id"]))
        if item and item.get("visto_em"):
            # plano já visado pela supervisão: vale a versão gravada (não aceita edições)
            dados, plano, visto = item["dados"], item["plano"], _visto_de(item)
        else:
            try:
                db.atualizar(int(body["id"]), dados, plano)   # guarda as edições feitas na tela
            except Exception:
                pass
    pdf = gerar_pdf(dados, plano, visto)
    nome = f"plano-de-aula-{_slug(dados.get('disciplina','')) or 'x'}-{_slug(dados.get('serie',''))}-{_slug(dados.get('data',''))}-{_slug(dados.get('professor',''))}.pdf"
    resp = send_file(BytesIO(pdf), mimetype="application/pdf", as_attachment=True, download_name=nome)
    # envio automático ao Drive (silencioso: nunca impede o download)
    if body.get("id") and drive.conectado() and drive.status()["auto"]:
        try:
            info = _enviar_drive(int(body["id"]), dados, plano, pdf)
            resp.headers["X-Drive-Link"] = info["link"]
        except Exception as e:  # noqa: BLE001
            resp.headers["X-Drive-Erro"] = str(e)[:150]
    return resp


def _enviar_drive(id_: int, dados: dict, plano: dict, pdf: bytes | None = None) -> dict:
    item = db.obter(id_) or {}
    if item.get("visto_em"):
        # plano visado: o Drive recebe sempre a versão assinada (com carimbo)
        dados, plano = item["dados"], item["plano"]
        pdf = gerar_pdf(dados, plano, _visto_de(item))
    pdf = pdf or gerar_pdf(dados, plano)
    d = dict(dados); d["_tema"] = plano.get("tema", "")
    em = (item.get("professor_email") or dados.get("professor_email") or "").strip()
    if em:
        u = db.usuario_por_email(em)
        if u and u.get("drive_pasta_id"):
            d["_pasta_id"] = u["drive_pasta_id"]
    info = drive.enviar_pdf(pdf, d, file_id_existente=item.get("drive_file_id"))
    db.set_drive(id_, info["id"], info["link"])
    return info


# ----------------------------------------------------------------------------
# Histórico
# ----------------------------------------------------------------------------
def _visto_de(item: dict) -> dict | None:
    if not item or not item.get("visto_em"):
        return None
    return {"em": item["visto_em"], "por": item.get("visto_por") or "", "codigo": item.get("visto_codigo") or ""}


def _codigo_visto(id_: int, quando: str) -> str:
    """Código curto de verificação (impresso no carimbo), derivado do plano + data/hora + chave secreta."""
    h = hashlib.sha256(f"{id_}|{quando}|{app.secret_key}".encode()).hexdigest().upper()
    return f"{h[:4]}-{h[4:8]}"


def _medir_banco() -> str:
    """Mede: (1) pegar conexão + 1ª consulta; (2) consultas seguintes na MESMA conexão (= latência pura da rede)."""
    import time as _t
    try:
        t0 = _t.time()
        with db.conexao() as con:
            con.execute("SELECT 1")
            ms1 = (_t.time() - t0) * 1000
            t1 = _t.time()
            for _ in range(3):
                con.execute("SELECT 1")
            ms2 = (_t.time() - t1) * 1000 / 3
        pool = "pool ativo" if getattr(db, "_POOL", None) else ("SEM pool" if db.USA_PG else "SQLite")
        host = ""
        if db.USA_PG:
            import re as _re
            m = _re.search(r"@([^/:?]+)", db.DATABASE_URL)
            host = m.group(1) if m else ""
        onde = ""
        if "neon.tech" in host:
            onde = " · Neon" + (" São Paulo" if "sa-east-1" in host else "") + (" (via pooler)" if "-pooler" in host else "")
        elif "render.com" in host:
            onde = " · Render"
        if ms2 > 150:
            nota = (" ℹ️ latência de rede entre o servidor (Render, EUA) e o banco (São Paulo); esperado, "
                    "o pool de conexões compensa a maior parte") if "sa-east-1" in host else \
                   " ⚠️ banco distante do servidor"
        else:
            nota = " ✅"
        return f"obter conexão + consulta: {ms1:.0f} ms · consulta na mesma conexão: {ms2:.0f} ms{nota} · {pool}{onde}"
    except Exception as e:  # noqa: BLE001
        return f"❌ {e}"


def _pode_ver(item: dict) -> bool:
    if _e_supervisao():
        return True
    em = _email_professor()
    if em and item.get("dados", {}).get("professor_email"):
        return item["dados"]["professor_email"].lower() == em.lower()
    return item.get("professor", "").lower() == _nome_professor().lower()


@app.route("/historico")
def historico():
    sup = _e_supervisao()
    return render_template(
        "historico.html",
        e_supervisao=sup,
        tem_senha_sup=bool(SUPERVISAO_SENHA),
        usuario=usuario_atual(),
        professor_atual=_nome_professor(),
        professores=db.professores() if sup else [],
        disciplinas=DISCIPLINAS, series=SERIES,
    )


@app.route("/api/historico")
def api_historico():
    sup = _e_supervisao()
    prof = request.args.get("professor", "").strip()
    if sup:
        itens = db.listar(professor=prof or None, disciplina=request.args.get("disciplina", ""),
                          serie=request.args.get("serie", ""), busca=request.args.get("busca", ""))
    else:
        itens = db.listar(professor=_nome_professor(), professor_email=_email_professor() or None,
                          disciplina=request.args.get("disciplina", ""),
                          serie=request.args.get("serie", ""), busca=request.args.get("busca", ""))
    return jsonify(itens)


@app.route("/api/historico/<int:id_>")
def api_historico_item(id_):
    item = db.obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver(item):
        return jsonify({"erro": "sem permissão"}), 403
    return jsonify(item)


@app.route("/api/historico/<int:id_>/pdf")
def api_historico_pdf(id_):
    item = db.obter(id_)
    if not item:
        return "não encontrado", 404
    if not _pode_ver(item):
        return "sem permissão", 403
    pdf = gerar_pdf(item["dados"], item["plano"], _visto_de(item))
    d = item["dados"]
    nome = f"plano-{_slug(d.get('disciplina',''))}-{_slug(d.get('serie',''))}-{_slug(d.get('data',''))}-{_slug(d.get('professor',''))}.pdf"
    return send_file(BytesIO(pdf), mimetype="application/pdf", as_attachment=True, download_name=nome)


@app.route("/api/historico/<int:id_>", methods=["DELETE"])
def api_historico_excluir(id_):
    item = db.obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver(item):
        return jsonify({"erro": "sem permissão"}), 403
    db.excluir(id_)
    return jsonify({"ok": True})


def _dar_visto(item: dict, usuario: dict | None = None, base_url: str | None = None) -> dict:
    """Aplica o visto eletrônico em um plano: grava, atualiza Drive e avisa o professor. Devolve {ok, codigo, aviso}.
    (usuario e base_url são passados explicitamente quando chamado fora da requisição, ex.: em threads)"""
    id_ = item["id"]
    u = usuario if usuario is not None else (usuario_atual() or {})
    base_url = base_url or PUBLIC_URL or _base_url()
    por = (u.get("nome") or "").strip() or SUPERVISAO_NOME or "Supervisão Pedagógica"
    quando = fuso.agora_txt()
    codigo = _codigo_visto(id_, quando)
    with db.conexao() as con:
        con.execute(db._q("UPDATE planos SET visto_em=?, visto_por=?, visto_email=?, visto_codigo=? WHERE id=?"),
                    (quando, por, u.get("email") or None, codigo, id_))
    if not (item["dados"].get("supervisao") or "").strip():
        item["dados"]["supervisao"] = por
        db.atualizar(id_, item["dados"], item["plano"])
    aviso = ""
    # Drive: quem envia é o professor. Aqui só atualizamos o arquivo que JÁ está no Drive (para receber o carimbo),
    # e apenas se a opção "atualizar no Drive ao dar visto" estiver ligada (padrão: ligada).
    try:
        if item.get("drive_file_id") and drive.conectado() and db.config_get("drive_visto_atualiza", "1") == "1":
            _enviar_drive(id_, item["dados"], item["plano"])
    except Exception as e:  # noqa: BLE001
        aviso = f"não foi possível atualizar o PDF no Drive ({e}). "
    dest = (item.get("professor_email") or item["dados"].get("professor_email") or "").strip()
    if dest and email_util.configurado():
        try:
            d = item["dados"]
            link = base_url + f"/?id={id_}"
            dt = datetime.strptime(quando[:19], "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y às %H:%M")
            txt = (f"Olá, {d.get('professor','')}!\n\nSeu plano de aula de {d.get('disciplina','')} ({d.get('serie','')}), "
                   f"semana {d.get('data','')}, tema \"{item['plano'].get('tema','')}\", recebeu o visto eletrônico da supervisão.\n\n"
                   f"Visto por: {por}\nEm: {dt}\nCódigo: {codigo}\n\n"
                   f"Baixe a versão assinada em: {link}\n\n"
                   "Observação: após o visto o plano fica travado para edição. Se precisar alterar, peça à supervisão para desfazer o visto.\n\n"
                   "Escola Estadual Presidente Bernardes — Assistente de Plano de Aula")
            html = email_util.template_simples("Plano de aula visado ✔", d.get("professor", ""),
                                               f"Seu plano de <b>{d.get('disciplina','')}</b> ({d.get('serie','')}), semana {d.get('data','')}, "
                                               f"tema <i>{item['plano'].get('tema','')}</i>, recebeu o <b>visto eletrônico da supervisão</b>.<br><br>"
                                               f"Visto por: <b>{por}</b><br>Em: {dt}<br>Código: {codigo}",
                                               link, "Abrir plano assinado")
            email_util.enviar(dest, f"Plano de aula visado – {d.get('disciplina','')} – {d.get('data','')}", txt, html)
        except Exception as e:  # noqa: BLE001
            aviso += f"o e-mail ao professor falhou ({e})."
    elif not dest:
        aviso += "professor sem e-mail cadastrado neste plano, não foi avisado."
    return {"ok": True, "visto_em": quando, "visto_por": por, "codigo": codigo, "aviso": aviso}


@app.route("/api/historico/<int:id_>/visto", methods=["POST"])
def api_historico_visto(id_):
    """Visto eletrônico da supervisão em um plano."""
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    item = db.obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    body = request.get_json(silent=True) or {}
    if body.get("desfazer"):
        db.marcar_visto(id_, "", desfazer=True)
        return jsonify({"ok": True})
    r = _dar_visto(item)
    if r["aviso"]:
        r["aviso"] = "Visto registrado, mas " + r["aviso"]
    return jsonify(r)


@app.route("/api/historico/visto-em-bloco", methods=["POST"])
def api_historico_visto_bloco():
    """Assina vários planos de uma vez. body: {ids: [..]}.
    Processa no máximo 5 por chamada, em paralelo (o navegador repete até terminar) — evita estourar o tempo do servidor."""
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    body = request.get_json(silent=True) or {}
    ids = [int(x) for x in body.get("ids", []) if str(x).isdigit()][:5]
    u = usuario_atual() or {}
    base_url = PUBLIC_URL or _base_url()
    itens = []
    pulados = 0
    for id_ in ids:
        item = db.obter(id_)
        if not item or item.get("visto_em"):
            pulados += 1
        else:
            itens.append(item)
    feitos, avisos = 0, []
    from concurrent.futures import ThreadPoolExecutor

    def trabalho(item):
        return item, _dar_visto(item, usuario=u, base_url=base_url)

    with ThreadPoolExecutor(max_workers=5) as ex:
        for fut in [ex.submit(trabalho, it) for it in itens]:
            try:
                item, r = fut.result()
                feitos += 1
                if r["aviso"]:
                    avisos.append(f"#{item['id']} {item['professor']}: {r['aviso']}")
            except Exception as e:  # noqa: BLE001
                avisos.append(f"falhou ({e})")
    return jsonify({"ok": True, "feitos": feitos, "pulados": pulados, "avisos": avisos})


def _planos_visiveis(limite: int = 2000) -> list[dict]:
    """Planos que o usuário atual pode ver (todos para supervisão; só os seus para professor)."""
    if _e_supervisao():
        return db.listar(limite=limite)
    return db.listar(professor=_nome_professor(), professor_email=_email_professor() or None, limite=limite)


@app.route("/api/historico/drive-em-bloco", methods=["POST"])
def api_historico_drive_bloco():
    """Envia ao Drive todos os planos ainda não enviados: o professor envia os dele; a supervisão, todos."""
    if not drive.conectado():
        return jsonify({"erro": "Google Drive não conectado. Peça à supervisão para conectar em Configurações."}), 400
    body = request.get_json(silent=True) or {}
    permitidos = {i["id"] for i in _planos_visiveis()}
    ids = [int(x) for x in body.get("ids", []) if str(x).isdigit() and int(x) in permitidos]
    if not ids:
        ids = [i["id"] for i in _planos_visiveis() if not i.get("drive_link")]
    ids = ids[:5]   # lote pequeno por chamada (o navegador repete até acabar)
    feitos, erros = 0, []
    from concurrent.futures import ThreadPoolExecutor

    def trabalho(id_):
        item = db.obter(id_)
        if item:
            _enviar_drive(id_, item["dados"], item["plano"])
        return item

    with ThreadPoolExecutor(max_workers=5) as ex:
        futs = {ex.submit(trabalho, i): i for i in ids}
        for fut, id_ in futs.items():
            try:
                if fut.result():
                    feitos += 1
            except Exception as e:  # noqa: BLE001
                erros.append(f"#{id_}: {e}")
    restantes = len([i for i in _planos_visiveis() if not i.get("drive_link")])
    return jsonify({"ok": True, "feitos": feitos, "erros": erros, "restantes": restantes})


# ----------------------------------------------------------------------------
# Google Drive
# ----------------------------------------------------------------------------
def _redirect_uri() -> str:
    base = os.getenv("PUBLIC_URL", "").strip().rstrip("/") or request.url_root.rstrip("/")
    if base.startswith("http://") and "localhost" not in base and "127.0.0.1" not in base:
        base = "https://" + base[len("http://"):]
    return base + "/drive/callback"


@app.route("/api/drive/status")
def api_drive_status():
    st = drive.status()
    st["redirect_uri"] = _redirect_uri()
    return jsonify(st)


@app.route("/drive/conectar")
def drive_conectar():
    if not _e_supervisao():
        return redirect(url_for("sup_login"))
    if not drive.credenciais_ok():
        return "Defina GOOGLE_CLIENT_ID e GOOGLE_CLIENT_SECRET no servidor.", 400
    state = os.urandom(12).hex()
    session["drive_state"] = state
    return redirect(drive.url_autorizacao(_redirect_uri(), state))


@app.route("/drive/callback")
def drive_callback():
    if request.args.get("error"):
        return redirect("/?drive=erro&msg=" + request.args["error"])
    if request.args.get("state") != session.get("drive_state"):
        return "Estado inválido. Tente conectar novamente.", 400
    try:
        drive.trocar_codigo(request.args.get("code", ""), _redirect_uri())
        drive.pasta_raiz()
    except Exception as e:  # noqa: BLE001
        return f"Falha ao conectar ao Google Drive: {e}", 500
    return redirect("/?drive=ok")


@app.route("/api/drive/desconectar", methods=["POST"])
def api_drive_desconectar():
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    drive.desconectar()
    return jsonify({"ok": True})


@app.route("/api/drive/config", methods=["POST"])
def api_drive_config():
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    body = request.get_json(force=True) or {}
    try:
        if "auto" in body:
            db.config_set("drive_auto", "1" if body["auto"] else "0")
        if body.get("estrutura") in ("professor", "serie"):
            db.config_set("drive_estrutura", body["estrutura"])
        if "visto_atualiza" in body:
            db.config_set("drive_visto_atualiza", "1" if body["visto_atualiza"] else "0")
        if body.get("pasta") is not None and drive.conectado():
            drive.definir_pasta_raiz(body["pasta"])
    except Exception as e:  # noqa: BLE001
        return jsonify({"erro": str(e)}), 400
    return jsonify(drive.status())


@app.route("/api/historico/<int:id_>/drive", methods=["POST"])
def api_historico_drive(id_):
    """Envia (ou reenvia) um plano do histórico para o Drive."""
    item = db.obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver(item):
        return jsonify({"erro": "sem permissão"}), 403
    if not drive.conectado():
        return jsonify({"erro": "Google Drive não conectado. Peça à supervisão para conectar em Configurações."}), 400
    try:
        info = _enviar_drive(id_, item["dados"], item["plano"])
    except Exception as e:  # noqa: BLE001
        return jsonify({"erro": f"Falha ao enviar ao Drive: {e}"}), 502
    return jsonify(info)


# ----------------------------------------------------------------------------
# Usuários (supervisão)
# ----------------------------------------------------------------------------
# ----------------------------------------------------------------------------
# Provas e atividades (a partir dos planos de aula)
# ----------------------------------------------------------------------------
def _pode_ver_atividade(item: dict) -> bool:
    if _e_supervisao():
        return True
    em = _email_professor()
    if em and item.get("professor_email"):
        return item["professor_email"].lower() == em.lower()
    return (item.get("professor") or "").lower() == _nome_professor().lower()


def _meta_atividade(item_params: dict, conteudo: dict) -> dict:
    return {
        "disciplina": item_params.get("disciplina", ""), "serie": item_params.get("serie", ""),
        "professor": item_params.get("professor", ""),
        "tipo_nome": TIPOS_ATIVIDADE.get(item_params.get("tipo", "prova"), "Prova"),
        "avaliativa": bool(item_params.get("avaliativa")), "valor_total": item_params.get("valor_total"),
    }


def _params_limpos(p: dict) -> dict:
    """Normaliza os parâmetros vindos do navegador."""
    out = {
        "tipo": p.get("tipo") if p.get("tipo") in TIPOS_ATIVIDADE else "prova",
        "disciplina": str(p.get("disciplina", "")).strip(), "serie": str(p.get("serie", "")).strip(),
        "n_me": max(0, min(30, int(p.get("n_me") or 0))), "n_disc": max(0, min(30, int(p.get("n_disc") or 0))),
        "dificuldade": p.get("dificuldade") or "mista", "observacoes": str(p.get("observacoes", ""))[:800],
        "avaliativa": bool(p.get("avaliativa")),
        "planos_ids": [int(x) for x in (p.get("planos_ids") or []) if str(x).isdigit()][:20],
    }
    if out["avaliativa"]:
        try:
            out["valor_total"] = float(str(p.get("valor_total") or 10).replace(",", "."))
        except ValueError:
            out["valor_total"] = 10.0
    return out


@app.route("/atividades")
def atividades():
    return render_template("atividades.html", e_supervisao=_e_supervisao(), usuario=usuario_atual(),
                           professor_atual=_nome_professor(), tipos=TIPOS_ATIVIDADE,
                           disciplinas=DISCIPLINAS, series=SERIES, ia_configurada=bool(carregar_config()["api_key"]))


@app.route("/api/atividades/gerar", methods=["POST"])
def api_atividades_gerar():
    p = _params_limpos(request.get_json(force=True) or {})
    if not p["planos_ids"]:
        return jsonify({"erro": "Selecione pelo menos um plano de aula."}), 400
    if p["n_me"] + p["n_disc"] <= 0:
        return jsonify({"erro": "Informe a quantidade de questões."}), 400
    if p["n_me"] + p["n_disc"] > 30:
        return jsonify({"erro": "Máximo de 30 questões por prova/atividade."}), 400
    planos = [pl for pl in db.planos_obter_varios(p["planos_ids"]) if _pode_ver(pl)]
    if not planos:
        return jsonify({"erro": "Planos não encontrados."}), 404
    p["planos_ids"] = [pl["id"] for pl in planos]
    p["disciplina"] = p["disciplina"] or planos[0]["disciplina"]
    p["serie"] = p["serie"] or planos[0]["serie"]
    u = usuario_atual()
    p["professor"] = (u or {}).get("nome") or planos[0]["professor"]
    p["professor_email"] = (u or {}).get("email", "")
    try:
        conteudo = gerar_atividade(p, planos)
    except ValueError as e:
        return jsonify({"erro": str(e)}), 400
    except Exception as e:  # noqa: BLE001
        msg = re.sub(r"[?&]key=[^&\s)\]]+", "?key=***", str(e))
        return jsonify({"erro": f"Falha ao consultar a IA: {msg}"}), 502
    id_ = db.atividade_inserir(p, conteudo)
    return jsonify({"id": id_, "params": p, "conteudo": conteudo})


@app.route("/api/atividades")
def api_atividades():
    a = request.args
    if _e_supervisao():
        itens = db.atividades_listar(professor=a.get("professor", "").strip() or None, disciplina=a.get("disciplina", ""),
                                     serie=a.get("serie", ""), busca=a.get("busca", ""))
    else:
        itens = db.atividades_listar(professor=_nome_professor(), professor_email=_email_professor() or None,
                                     disciplina=a.get("disciplina", ""), serie=a.get("serie", ""), busca=a.get("busca", ""))
    return jsonify(itens)


@app.route("/api/atividades/<int:id_>")
def api_atividade_item(id_):
    item = db.atividade_obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver_atividade(item):
        return jsonify({"erro": "sem permissão"}), 403
    return jsonify(item)


@app.route("/api/atividades/<int:id_>/salvar", methods=["POST"])
def api_atividade_salvar(id_):
    item = db.atividade_obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver_atividade(item):
        return jsonify({"erro": "sem permissão"}), 403
    body = request.get_json(force=True) or {}
    params = dict(item["params"])
    novos = body.get("params") or {}
    for k in ("tipo", "avaliativa", "valor_total"):
        if k in novos:
            params[k] = novos[k]
    params["avaliativa"] = bool(params.get("avaliativa"))
    conteudo = body.get("conteudo") or item["conteudo"]
    qs = []
    for q in conteudo.get("questoes") or []:
        q = {k: v for k, v in q.items() if not str(k).startswith("_")}
        if params["avaliativa"]:
            try:
                q["valor"] = round(float(str(q.get("valor", 0)).replace(",", ".")), 2)
            except ValueError:
                q["valor"] = 0
        else:
            q.pop("valor", None)
        qs.append(q)
    conteudo["questoes"] = qs
    if params["avaliativa"]:
        params["valor_total"] = round(sum(float(q.get("valor") or 0) for q in qs), 2)
    db.atividade_atualizar(id_, params, conteudo)
    aviso = ""
    if item.get("drive_file_id") and drive.conectado():
        try:
            _enviar_drive_atividade(id_, params, conteudo)
        except Exception as e:  # noqa: BLE001
            aviso = f"Salvo, mas não foi possível atualizar o PDF no Drive: {str(e)[:120]}"
    return jsonify({"ok": True, "params": params, "aviso": aviso})


@app.route("/api/atividades/<int:id_>/pdf")
def api_atividade_pdf(id_):
    item = db.atividade_obter(id_)
    if not item:
        return "não encontrado", 404
    if not _pode_ver_atividade(item):
        return "sem permissão", 403
    gabarito = request.args.get("gabarito") == "1"
    meta = _meta_atividade(item["params"], item["conteudo"])
    pdf = (gerar_pdf_gabarito if gabarito else gerar_pdf_atividade)(item["conteudo"], meta)
    p = item["params"]
    nome = f"{'gabarito' if gabarito else _slug(p.get('tipo','prova'))}-{_slug(p.get('disciplina',''))}-{_slug(p.get('serie',''))}-{_slug(item['conteudo'].get('titulo',''))[:40]}.pdf"
    return send_file(BytesIO(pdf), mimetype="application/pdf", as_attachment=True, download_name=nome)


@app.route("/api/atividades/<int:id_>", methods=["DELETE"])
def api_atividade_excluir(id_):
    item = db.atividade_obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver_atividade(item):
        return jsonify({"erro": "sem permissão"}), 403
    db.atividade_excluir(id_)
    return jsonify({"ok": True})


def _enviar_drive_atividade(id_: int, params: dict, conteudo: dict) -> dict:
    item = db.atividade_obter(id_) or {}
    meta = _meta_atividade(params, conteudo)
    pdf = gerar_pdf_atividade(conteudo, meta)
    d = {"professor": params.get("professor", ""), "disciplina": params.get("disciplina", ""), "serie": params.get("serie", ""),
         "data": "", "_tema": conteudo.get("titulo", ""),
         "_nome_arquivo": f"{meta['tipo_nome']} - {params.get('disciplina','')} - {params.get('serie','')} - {conteudo.get('titulo','')[:60]} - {params.get('professor','')}.pdf"}
    em = (item.get("professor_email") or params.get("professor_email") or "").strip()
    if em:
        u = db.usuario_por_email(em)
        if u and u.get("drive_pasta_id"):
            d["_pasta_id"] = u["drive_pasta_id"]
    info = drive.enviar_pdf(pdf, d, file_id_existente=item.get("drive_file_id"))
    db.atividade_set_drive(id_, info["id"], info["link"])
    return info


@app.route("/api/atividades/<int:id_>/drive", methods=["POST"])
def api_atividade_drive(id_):
    item = db.atividade_obter(id_)
    if not item:
        return jsonify({"erro": "não encontrado"}), 404
    if not _pode_ver_atividade(item):
        return jsonify({"erro": "sem permissão"}), 403
    if not drive.conectado():
        return jsonify({"erro": "Google Drive não conectado."}), 400
    try:
        info = _enviar_drive_atividade(id_, item["params"], item["conteudo"])
    except Exception as e:  # noqa: BLE001
        return jsonify({"erro": f"Falha ao enviar ao Drive: {str(e)[:200]}"}), 502
    return jsonify(info)

# ----------------------------------------------------------------------------
# Correção: aplicações (prova online / lançamento em papel), respostas, relatório
# ----------------------------------------------------------------------------
import random as _random  # noqa: E402
import string as _string  # noqa: E402

_ALFA = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"   # sem 0/O/1/I para não confundir no quadro


def _novo_codigo() -> str:
    for _ in range(20):
        c = "".join(_random.choice(_ALFA) for _ in range(6))
        if not db.aplicacao_por_codigo(c):
            return c
    return "".join(_random.choice(_ALFA) for _ in range(8))


def _atividade_e_dono(id_: int):
    item = db.atividade_obter(id_)
    if not item:
        return None, (jsonify({"erro": "não encontrada"}), 404)
    if not _pode_ver_atividade(item):
        return None, (jsonify({"erro": "sem permissão"}), 403)
    return item, None


def _aplicacao_e_dono(ap_id: int):
    ap = db.aplicacao_obter(ap_id)
    if not ap:
        return None, None, (jsonify({"erro": "aplicação não encontrada"}), 404)
    item, err = _atividade_e_dono(ap["atividade_id"])
    if err:
        return None, None, err
    return ap, item, None


@app.route("/api/atividades/<int:id_>/aplicacoes")
def api_aplicacoes(id_):
    item, err = _atividade_e_dono(id_)
    if err:
        return err
    aps = db.aplicacoes_da_atividade(id_)
    for a in aps:
        a["link"] = f"{_base_url()}/prova/{a['codigo']}"
    return jsonify(aps)


@app.route("/api/atividades/<int:id_>/aplicacoes", methods=["POST"])
def api_aplicacao_criar(id_):
    item, err = _atividade_e_dono(id_)
    if err:
        return err
    b = request.get_json(force=True) or {}
    tempo = b.get("tempo_min")
    try:
        tempo = int(tempo) if tempo else None
    except ValueError:
        tempo = None
    ap_id = db.aplicacao_criar(id_, _novo_codigo(), str(b.get("turma") or item["params"].get("serie") or "")[:60],
                               embaralhar=bool(b.get("embaralhar", True)), mostrar_nota=bool(b.get("mostrar_nota", False)),
                               tempo_min=tempo)
    ap = db.aplicacao_obter(ap_id)
    ap["link"] = f"{_base_url()}/prova/{ap['codigo']}"
    return jsonify(ap)


@app.route("/api/aplicacoes/<int:ap_id>", methods=["PATCH"])
def api_aplicacao_editar(ap_id):
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    b = request.get_json(force=True) or {}
    campos = {}
    if "aberta" in b:
        campos["aberta"] = 1 if b["aberta"] else 0
        campos["encerrada_em"] = None if b["aberta"] else fuso.agora_txt()
    for k in ("mostrar_nota", "embaralhar"):
        if k in b:
            campos[k] = 1 if b[k] else 0
    if "turma" in b:
        campos["turma"] = str(b["turma"])[:60]
    db.aplicacao_atualizar(ap_id, **campos)
    return jsonify({"ok": True})


@app.route("/api/aplicacoes/<int:ap_id>", methods=["DELETE"])
def api_aplicacao_excluir(ap_id):
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    db.aplicacao_excluir(ap_id)
    return jsonify({"ok": True})


@app.route("/api/aplicacoes/<int:ap_id>/qr")
def api_aplicacao_qr(ap_id):
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    import segno
    link = f"{_base_url()}/prova/{ap['codigo']}"
    buf = BytesIO()
    segno.make(link, error="m").save(buf, kind="png", scale=8, border=2)
    buf.seek(0)
    return send_file(buf, mimetype="image/png")


@app.route("/api/aplicacoes/<int:ap_id>/folha")
def api_aplicacao_folha(ap_id):
    """PDF para projetar/imprimir: título, código, link e QR code."""
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    from pdf_atividade import gerar_pdf_folha_acesso
    link = f"{_base_url()}/prova/{ap['codigo']}"
    pdf = gerar_pdf_folha_acesso(item["conteudo"].get("titulo", ""), _meta_atividade(item["params"], item["conteudo"]),
                                 ap["codigo"], link, ap.get("turma") or "")
    return send_file(BytesIO(pdf), mimetype="application/pdf", as_attachment=True,
                     download_name=f"acesso-prova-{ap['codigo']}.pdf")


def _questoes_para_aluno(item: dict, ap: dict, semente: str) -> list[dict]:
    """Questões sem gabarito/resolução. Embaralha a ordem das alternativas por aluno (mantendo índice original)."""
    qs = item["conteudo"].get("questoes") or []
    rnd = _random.Random(f"{ap['id']}|{semente}")
    out = []
    for i, q in enumerate(qs):
        item_q = {"i": i, "tipo": q.get("tipo"), "enunciado": q.get("enunciado", ""),
                  "valor": q.get("valor") if item["params"].get("avaliativa") else None}
        if q.get("tipo") == "me":
            alts = list(enumerate((q.get("alternativas") or [])[:5]))
            if ap.get("embaralhar"):
                rnd.shuffle(alts)
            item_q["alternativas"] = [{"k": k, "texto": t} for k, t in alts]
        else:
            item_q["linhas"] = q.get("linhas", 6)
        out.append(item_q)
    return out


@app.route("/prova", methods=["GET", "POST"])
def prova_entrar():
    """Página pública: aluno digita o código da prova."""
    erro = ""
    if request.method == "POST":
        cod = (request.form.get("codigo") or "").strip().upper().replace(" ", "")
        ap = db.aplicacao_por_codigo(cod) if cod else None
        if ap:
            return redirect(url_for("prova_aluno", codigo=ap["codigo"]))
        erro = "Código não encontrado. Confira com o professor."
    return render_template("prova_entrar.html", erro=erro)


@app.route("/prova/<codigo>")
def prova_aluno(codigo):
    ap = db.aplicacao_por_codigo(codigo)
    if not ap:
        return render_template("prova_entrar.html", erro="Código não encontrado. Confira com o professor."), 404
    item = db.atividade_obter(ap["atividade_id"])
    return render_template("prova_aluno.html", codigo=ap["codigo"], titulo=item["conteudo"].get("titulo", ""),
                           disciplina=item["params"].get("disciplina", ""), serie=ap.get("turma") or item["params"].get("serie", ""),
                           professor=item["params"].get("professor", ""), aberta=bool(ap["aberta"]),
                           tempo_min=ap.get("tempo_min"), instrucoes=item["conteudo"].get("instrucoes", ""),
                           avaliativa=bool(item["params"].get("avaliativa")), valor_total=item["params"].get("valor_total"))


@app.route("/api/prova/<codigo>/info", methods=["POST"])
def api_prova_aluno_info(codigo):
    """Aluno se identifica e recebe as questões (sem gabarito)."""
    ap = db.aplicacao_por_codigo(codigo)
    if not ap:
        return jsonify({"erro": "Código não encontrado."}), 404
    if not ap["aberta"]:
        return jsonify({"erro": "Esta prova já foi encerrada pelo professor."}), 403
    b = request.get_json(force=True) or {}
    nome = str(b.get("nome") or "").strip()[:80]
    numero = str(b.get("numero") or "").strip()[:10]
    if len(nome) < 3:
        return jsonify({"erro": "Digite seu nome completo."}), 400
    ja = db.resposta_ja_enviada(ap["id"], nome, numero)
    if ja:
        return jsonify({"erro": f"Já existe uma prova enviada por {ja['aluno_nome']} (nº {ja['aluno_numero'] or '-'}) em {ja['enviado_em'][11:16]}. Fale com o professor."}), 409
    item = db.atividade_obter(ap["atividade_id"])
    return jsonify({"questoes": _questoes_para_aluno(item, ap, f"{nome}|{numero}".lower()), "tempo_min": ap.get("tempo_min")})


@app.route("/api/prova/<codigo>/enviar", methods=["POST"])
def api_prova_aluno_enviar(codigo):
    ap = db.aplicacao_por_codigo(codigo)
    if not ap:
        return jsonify({"erro": "Código não encontrado."}), 404
    if not ap["aberta"]:
        return jsonify({"erro": "Esta prova já foi encerrada pelo professor."}), 403
    b = request.get_json(force=True) or {}
    nome = str(b.get("nome") or "").strip()[:80]
    numero = str(b.get("numero") or "").strip()[:10]
    if len(nome) < 3:
        return jsonify({"erro": "Nome inválido."}), 400
    if db.resposta_ja_enviada(ap["id"], nome, numero):
        return jsonify({"erro": "Sua prova já havia sido enviada."}), 409
    respostas = {str(k): v for k, v in (b.get("respostas") or {}).items()}
    item = db.atividade_obter(ap["atividade_id"])
    n_q = len(item["conteudo"].get("questoes") or [])
    respostas = {k: (v if isinstance(v, (int, str)) else "") for k, v in respostas.items() if k.isdigit() and int(k) < n_q}
    for k, v in list(respostas.items()):
        if isinstance(v, str):
            respostas[k] = v[:4000]
    rid = db.resposta_criar(ap["id"], nome, numero, respostas, origem="online")
    # corrige ME na hora; discursivas com IA (em seguida, síncrono — poucos segundos)
    corr = correcao.corrigir(item, respostas, usar_ia=True)
    db.resposta_corrigir(rid, corr, corr["nota"], corr["nota_me"], corr["nota_disc"], corr["status"])
    resp = {"ok": True}
    if ap.get("mostrar_nota"):
        resp.update(acertos_me=corr["acertos_me"], total_me=corr["total_me"], total_disc=corr["total_disc"],
                    nota_me=corr["nota_me"], possivel=corr["possivel"],
                    nota=corr["nota"] if corr["status"] == "corrigido" and not corr["total_disc"] else None)
    return jsonify(resp)


@app.route("/api/aplicacoes/<int:ap_id>/respostas")
def api_respostas(ap_id):
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    rs = db.respostas_da_aplicacao(ap_id)
    return jsonify({"aplicacao": ap, "respostas": rs, "questoes": item["conteudo"].get("questoes") or [],
                    "avaliativa": bool(item["params"].get("avaliativa")), "estatisticas": correcao.estatisticas(item, rs)})


@app.route("/api/aplicacoes/<int:ap_id>/lancar", methods=["POST"])
def api_lancar_papel(ap_id):
    """Prova em papel: professor digita as respostas. body: {alunos:[{nome, numero, me:"BCADE...", disc:{i:texto}}]}"""
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    b = request.get_json(force=True) or {}
    qs = item["conteudo"].get("questoes") or []
    idx_me = [i for i, q in enumerate(qs) if q.get("tipo") == "me"]
    feitos, erros = 0, []
    for a in (b.get("alunos") or [])[:80]:
        nome = str(a.get("nome") or "").strip()[:80]
        if len(nome) < 2:
            continue
        numero = str(a.get("numero") or "").strip()[:10]
        letras = re.sub(r"[^A-Ea-e\-_.xX*]", "", str(a.get("me") or "")).upper()
        respostas = {}
        for pos, i in enumerate(idx_me):
            if pos < len(letras) and letras[pos] in "ABCDE":
                respostas[str(i)] = "ABCDE".index(letras[pos])
        for k, v in (a.get("disc") or {}).items():
            if str(k).isdigit():
                respostas[str(k)] = str(v)[:4000]
        notas_disc = a.get("notas_disc") or {}   # professor já corrigiu no papel: {i: nota}
        ja = db.resposta_ja_enviada(ap_id, nome, numero)
        if ja:
            db.resposta_atualizar(ja["id"], aluno_nome=nome, aluno_numero=numero, respostas=respostas)
            rid = ja["id"]
        else:
            rid = db.resposta_criar(ap_id, nome, numero, respostas, origem="papel")
        usar_ia = bool(b.get("usar_ia", True)) and not notas_disc
        corr = correcao.corrigir(item, respostas, usar_ia=usar_ia)
        for k, v in notas_disc.items():
            try:
                corr = correcao.aplicar_ajuste_professor(corr, int(k), float(str(v).replace(",", ".")), None, qs,
                                                         bool(item["params"].get("avaliativa")))
            except (ValueError, TypeError):
                pass
        if notas_disc or not [q for q in qs if q.get("tipo") == "disc"]:
            corr["status"] = "corrigido"
        db.resposta_corrigir(rid, corr, corr["nota"], corr["nota_me"], corr["nota_disc"], corr["status"])
        feitos += 1
    return jsonify({"ok": True, "feitos": feitos, "erros": erros})


@app.route("/api/respostas/<int:rid>/nota", methods=["POST"])
def api_resposta_nota(rid):
    """Professor confirma/ajusta a nota de uma questão. body: {i, nota, comentario}"""
    r = db.resposta_obter(rid)
    if not r:
        return jsonify({"erro": "não encontrada"}), 404
    ap, item, err = _aplicacao_e_dono(r["aplicacao_id"])
    if err:
        return err
    b = request.get_json(force=True) or {}
    corr = r["correcao"] or correcao.corrigir(item, r["respostas"], usar_ia=False)
    try:
        corr = correcao.aplicar_ajuste_professor(corr, int(b.get("i")), float(str(b.get("nota", 0)).replace(",", ".")),
                                                 b.get("comentario"), item["conteudo"].get("questoes") or [],
                                                 bool(item["params"].get("avaliativa")))
    except (ValueError, TypeError):
        return jsonify({"erro": "nota inválida"}), 400
    db.resposta_corrigir(rid, corr, corr["nota"], corr["nota_me"], corr["nota_disc"], corr["status"])
    return jsonify({"ok": True, "correcao": corr})


@app.route("/api/respostas/<int:rid>/confirmar", methods=["POST"])
def api_resposta_confirmar(rid):
    """Aceita todas as sugestões da IA desta prova."""
    r = db.resposta_obter(rid)
    if not r:
        return jsonify({"erro": "não encontrada"}), 404
    ap, item, err = _aplicacao_e_dono(r["aplicacao_id"])
    if err:
        return err
    corr = r["correcao"] or correcao.corrigir(item, r["respostas"], usar_ia=False)
    for it in corr["itens"]:
        if it["tipo"] == "disc":
            if it.get("nota") is None:
                it["nota"] = 0.0
            it["origem"] = "professor"; it["confianca"] = "alta"
    corr = correcao.aplicar_ajuste_professor(corr, -1, 0, None, item["conteudo"].get("questoes") or [], bool(item["params"].get("avaliativa")))
    corr["status"] = "corrigido"
    db.resposta_corrigir(rid, corr, corr["nota"], corr["nota_me"], corr["nota_disc"], corr["status"])
    return jsonify({"ok": True, "correcao": corr})


@app.route("/api/respostas/<int:rid>/recorrigir", methods=["POST"])
def api_resposta_recorrigir(rid):
    """Refaz a correção (ex.: após anular questão ou mudar gabarito). Mantém notas dadas pelo professor."""
    r = db.resposta_obter(rid)
    if not r:
        return jsonify({"erro": "não encontrada"}), 404
    ap, item, err = _aplicacao_e_dono(r["aplicacao_id"])
    if err:
        return err
    ant = {it["i"]: it for it in (r["correcao"] or {}).get("itens", [])}
    corr = correcao.corrigir(item, r["respostas"], usar_ia=bool((request.get_json(silent=True) or {}).get("usar_ia", True)), anteriores=ant)
    db.resposta_corrigir(rid, corr, corr["nota"], corr["nota_me"], corr["nota_disc"], corr["status"])
    return jsonify({"ok": True, "correcao": corr})


@app.route("/api/respostas/<int:rid>", methods=["DELETE"])
def api_resposta_excluir(rid):
    r = db.resposta_obter(rid)
    if not r:
        return jsonify({"erro": "não encontrada"}), 404
    ap, item, err = _aplicacao_e_dono(r["aplicacao_id"])
    if err:
        return err
    db.resposta_excluir(rid)
    return jsonify({"ok": True})


@app.route("/api/aplicacoes/<int:ap_id>/exportar")
def api_aplicacao_exportar(ap_id):
    """Excel com notas por aluno, acertos por questão e estatísticas."""
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rs = db.respostas_da_aplicacao(ap_id)
    qs = item["conteudo"].get("questoes") or []
    aval = bool(item["params"].get("avaliativa"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    cab = ["Nº", "Estudante", "Nota", "Múltipla escolha", "Discursivas", "Acertos ME", "Situação", "Enviado em", "Origem"] + \
          [f"Q{i+1}" for i in range(len(qs))]
    ws.append(cab)
    azul = PatternFill("solid", fgColor="1F4E79")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = azul; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sit = {"corrigido": "Corrigida", "revisar": "Revisar", "pendente": "Pendente"}
    for r in rs:
        corr = r.get("correcao") or {}
        linha = [r.get("aluno_numero") or "", r.get("aluno_nome"), r.get("nota"), r.get("nota_me"), r.get("nota_disc"),
                 f"{corr.get('acertos_me', '')}/{corr.get('total_me', '')}" if corr else "", sit.get(r.get("status"), r.get("status")),
                 _dt_br(r.get("enviado_em")), "Online" if r.get("origem") == "online" else "Papel"]
        por_i = {it["i"]: it for it in corr.get("itens", [])}
        for i, q in enumerate(qs):
            it = por_i.get(i)
            if not it:
                linha.append("")
            elif q.get("tipo") == "me":
                linha.append(("ABCDE"[it["marcada"]] if it.get("marcada") is not None else "—") + (" ✔" if it["acertou"] else ""))
            else:
                linha.append(it.get("nota") if it.get("nota") is not None else "?")
        ws.append(linha)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([6, 34, 8, 12, 12, 11, 11, 17, 8] + [7] * len(qs), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # aba questões
    est = correcao.estatisticas(item, rs)
    w2 = wb.create_sheet("Por questão")
    w2.append(["Questão", "Tipo", "Habilidade", "% acerto", "Gabarito", "A", "B", "C", "D", "E", "Enunciado"])
    for c in w2[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = azul
    for x in est["por_questao"]:
        d = x.get("distribuicao") or ["", "", "", "", ""]
        w2.append([x["i"] + 1, "ME" if x["tipo"] == "me" else "Disc.", x["habilidade"], x["pct"],
                   "ABCDE"[x["correta"]] if x["tipo"] == "me" and x.get("correta") is not None else "", *d, x["enunciado"]])
    for i, w in enumerate([8, 7, 14, 9, 9, 5, 5, 5, 5, 5, 80], 1):
        w2.column_dimensions[get_column_letter(i)].width = w
    w3 = wb.create_sheet("Resumo")
    for a, b in [("Prova", item["conteudo"].get("titulo", "")), ("Disciplina", item["params"].get("disciplina", "")),
                 ("Turma", ap.get("turma") or item["params"].get("serie", "")), ("Professor(a)", item["params"].get("professor", "")),
                 ("Código", ap["codigo"]), ("Respostas", est["n_respostas"]), ("Corrigidas", est["n_corrigidas"]),
                 ("Média", est["media"]), ("Maior nota", est["maior"]), ("Menor nota", est["menor"]),
                 ("Abaixo de 60%", est["abaixo_media"]), ("Faixas (<50 / 50-70 / 70-90 / ≥90)", " / ".join(map(str, est["faixas"])))]:
        w3.append([a, b])
    w3.column_dimensions["A"].width = 34; w3.column_dimensions["B"].width = 40
    for c in w3["A"]:
        c.font = Font(bold=True)
    if est["habilidades"]:
        w3.append([]); w3.append(["Habilidade", "% acerto", "Questões"])
        for h in est["habilidades"]:
            w3.append([h["habilidade"], h["pct"], ", ".join(map(str, h["questoes"]))])
    out = BytesIO(); wb.save(out); out.seek(0)
    return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True,
                     download_name=f"notas-{_slug(item['conteudo'].get('titulo',''))[:40]}-{ap['codigo']}.xlsx")


@app.route("/api/aplicacoes/<int:ap_id>/relatorio.pdf")
def api_aplicacao_relatorio(ap_id):
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return err
    from pdf_atividade import gerar_pdf_relatorio
    rs = db.respostas_da_aplicacao(ap_id)
    pdf = gerar_pdf_relatorio(item, ap, rs, correcao.estatisticas(item, rs))
    return send_file(BytesIO(pdf), mimetype="application/pdf", as_attachment=True,
                     download_name=f"resultado-{_slug(item['conteudo'].get('titulo',''))[:40]}-{ap['codigo']}.pdf")


@app.route("/correcao/<int:ap_id>")
def correcao_pagina(ap_id):
    ap, item, err = _aplicacao_e_dono(ap_id)
    if err:
        return redirect(url_for("atividades"))
    return render_template("correcao.html", ap=ap, item=item, e_supervisao=_e_supervisao(), usuario=usuario_atual(),
                           link=f"{_base_url()}/prova/{ap['codigo']}", tipos=TIPOS_ATIVIDADE)


@app.route("/usuarios")
def usuarios():
    if not _e_supervisao():
        return redirect(url_for("login"))
    return render_template("usuarios.html", e_supervisao=True, usuario=usuario_atual(),
                           dominio=auth.dominio_msg(), email_ok=email_util.configurado())


@app.route("/api/usuarios")
def api_usuarios():
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    return jsonify(db.usuarios_listar())


def _dt_br(txt: str) -> str:
    """'2026-09-15 14:03:00' -> '15/09/2026 14:03' (texto vazio se não houver)."""
    if not txt:
        return ""
    t = str(txt)[:16].replace("T", " ")
    try:
        return datetime.strptime(t, "%Y-%m-%d %H:%M").strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return t


@app.route("/usuarios/exportar")
def usuarios_exportar():
    """Baixa a lista de usuários em Excel (padrão) ou CSV (?formato=csv). Só supervisão."""
    if not _e_supervisao():
        return redirect(url_for("login"))
    formato = (request.args.get("formato") or "xlsx").lower()
    usuarios = db.usuarios_listar()
    try:
        stats = db.planos_por_professor()
    except Exception:
        stats = {}
    cab = ["Nome", "E-mail", "Perfil", "Situação", "Senha", "Planos criados", "Planos visados",
           "Último plano", "Último acesso", "Cadastrado em", "Pasta no Drive"]
    linhas = []
    for u in usuarios:
        st = stats.get((u.get("email") or "").lower()) or stats.get((u.get("nome") or "").strip().lower()) or {}
        linhas.append([
            u.get("nome") or "", u.get("email") or "",
            "Supervisão" if u.get("perfil") == "supervisao" else "Professor(a)",
            "Ativo" if u.get("ativo") else "Desativado",
            "Definida" if u.get("tem_senha") else "Pendente",
            st.get("total", 0), st.get("visados", 0), _dt_br(st.get("ultimo", "")),
            _dt_br(u.get("ultimo_acesso")), _dt_br(u.get("criado_em")),
            f"https://drive.google.com/drive/folders/{u['drive_pasta_id']}" if u.get("drive_pasta_id") else "",
        ])
    nome_arq = f"usuarios-{fuso.hoje().strftime('%Y-%m-%d')}"
    if formato == "csv":
        import csv
        from io import StringIO
        buf = StringIO()
        w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
        w.writerow(cab)
        w.writerows(linhas)
        dados = ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM: Excel abre com acentos certos
        return send_file(BytesIO(dados), mimetype="text/csv",
                         as_attachment=True, download_name=nome_arq + ".csv")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuários"
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(vertical="center")
    for l in linhas:
        ws.append(l)
    larguras = [32, 38, 13, 12, 10, 14, 14, 17, 17, 17, 60]
    for i, w_ in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w_
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    # aba resumo
    r = wb.create_sheet("Resumo")
    total = len(usuarios)
    ativos = sum(1 for u in usuarios if u.get("ativo"))
    pend = sum(1 for u in usuarios if u.get("ativo") and not u.get("tem_senha"))
    sup = sum(1 for u in usuarios if u.get("perfil") == "supervisao")
    planos = sum(l[5] for l in linhas)
    visados = sum(l[6] for l in linhas)
    for a, b in [("Gerado em", fuso.agora().strftime("%d/%m/%Y %H:%M")), ("Usuários cadastrados", total),
                 ("Ativos", ativos), ("Senha pendente", pend), ("Supervisão", sup),
                 ("Planos criados (total)", planos), ("Planos visados (total)", visados)]:
        r.append([a, b])
    r.column_dimensions["A"].width = 26
    r.column_dimensions["B"].width = 20
    for c in r["A"]:
        c.font = Font(bold=True)
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=nome_arq + ".xlsx")


@app.route("/api/usuarios", methods=["POST"])
def api_usuarios_criar():
    """Cadastra (um ou vários) e-mails e envia o link de criação de senha."""
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    body = request.get_json(force=True) or {}
    brutos = body.get("emails") or [body.get("email", "")]
    if isinstance(brutos, str):
        brutos = re.split(r"[\s,;]+", brutos)
    perfil = "supervisao" if body.get("perfil") == "supervisao" else "professor"
    enviar = body.get("enviar", True)
    resultados = []
    for e in brutos:
        e = auth.normalizar_email(e)
        if not e:
            continue
        if not auth.email_valido(e):
            resultados.append({"email": e, "status": "erro", "msg": f"fora do domínio {auth.dominio_msg()}"}); continue
        u = db.usuario_por_email(e)
        if not u:
            u = db.usuario_criar(e, body.get("nome", "") or auth.nome_do_email(e), perfil)
        elif perfil == "supervisao" and u["perfil"] != "supervisao":
            db.usuario_atualizar(u["id"], perfil="supervisao")
        if enviar:
            r = auth.solicitar_acesso(e, _base_url())
            resultados.append({"email": e, "status": "enviado" if r["enviado"] else "link",
                               "msg": "e-mail enviado" if r["enviado"] else (r["erro"] or ""), "link": r["link"]})
        else:
            resultados.append({"email": e, "status": "cadastrado", "msg": "cadastrado sem envio"})
    return jsonify(resultados)


@app.route("/api/usuarios/<int:id_>", methods=["PATCH"])
def api_usuarios_editar(id_):
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    body = request.get_json(force=True) or {}
    u = db.usuario_por_id(id_)
    if not u:
        return jsonify({"erro": "não encontrado"}), 404
    campos = {}
    if "nome" in body:
        campos["nome"] = body["nome"].strip()
    if "perfil" in body and body["perfil"] in ("professor", "supervisao"):
        campos["perfil"] = body["perfil"]
    if "ativo" in body:
        campos["ativo"] = 1 if body["ativo"] else 0
    if "drive_pasta" in body:
        v = (body["drive_pasta"] or "").strip()
        m = re.search(r"folders/([A-Za-z0-9_-]{10,})", v) or re.fullmatch(r"([A-Za-z0-9_-]{25,})", v)
        campos["drive_pasta_id"] = m.group(1) if m else None
    # impede remover a última supervisão
    if (campos.get("perfil") == "professor" or campos.get("ativo") == 0) and u["perfil"] == "supervisao" \
            and db.usuarios_total_supervisao() <= 1 and not SUPERVISAO_SENHA:
        return jsonify({"erro": "Não é possível remover a última conta de supervisão."}), 400
    db.usuario_atualizar(id_, **campos)
    return jsonify({"ok": True})


@app.route("/api/usuarios/<int:id_>", methods=["DELETE"])
def api_usuarios_excluir(id_):
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    u = db.usuario_por_id(id_)
    if not u:
        return jsonify({"erro": "não encontrado"}), 404
    if u["perfil"] == "supervisao" and db.usuarios_total_supervisao() <= 1 and not SUPERVISAO_SENHA:
        return jsonify({"erro": "Não é possível excluir a última conta de supervisão."}), 400
    db.usuario_excluir(id_)
    return jsonify({"ok": True})


@app.route("/api/email/teste", methods=["POST"])
def api_email_teste():
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    dest = (request.get_json(silent=True) or {}).get("para") or _email_professor()
    if not dest:
        return jsonify({"erro": "Informe o destinatário."}), 400
    try:
        email_util.enviar(dest, "Teste — Assistente de Plano de Aula",
                          "Este é um e-mail de teste. Se você o recebeu, o envio está funcionando.")
    except Exception as e:  # noqa: BLE001
        return jsonify({"erro": str(e)}), 502
    return jsonify({"ok": True, "modo": email_util.descricao()})


@app.route("/api/usuarios/<int:id_>/reenviar", methods=["POST"])
def api_usuarios_reenviar(id_):
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    u = db.usuario_por_id(id_)
    if not u:
        return jsonify({"erro": "não encontrado"}), 404
    r = auth.solicitar_acesso(u["email"], _base_url())
    return jsonify(r)


# ----------------------------------------------------------------------------
# Biblioteca de referências (documentos do ICE etc.)
# ----------------------------------------------------------------------------
@app.route("/biblioteca")
def biblioteca():
    return render_template("biblioteca.html", e_supervisao=_e_supervisao(), usuario=usuario_atual(),
                           fontes=referencias.FONTES)


@app.route("/api/documentos")
def api_documentos():
    return jsonify(db.documentos_listar(somente_ativos=not _e_supervisao()))


@app.route("/api/documentos", methods=["POST"])
def api_documentos_criar():
    if not _e_supervisao():
        return jsonify({"erro": "Apenas a supervisão pode enviar documentos."}), 403
    f = request.files.get("arquivo")
    if not f or not f.filename:
        return jsonify({"erro": "Selecione um arquivo PDF, DOCX ou TXT."}), 400
    nome = f.filename
    if not nome.lower().endswith((".pdf", ".docx", ".txt", ".md")):
        return jsonify({"erro": "Formato não suportado. Use PDF, DOCX ou TXT."}), 400
    conteudo = f.read()
    if len(conteudo) > 40 * 1024 * 1024:
        return jsonify({"erro": "Arquivo acima de 40 MB."}), 400
    try:
        texto = referencias.extrair_texto(nome, conteudo)
    except Exception as e:  # noqa: BLE001
        return jsonify({"erro": f"Não foi possível ler o arquivo: {e}"}), 400
    trechos = referencias.dividir_em_trechos(texto)
    if not trechos:
        return jsonify({"erro": "O arquivo não contém texto legível (pode ser um PDF só de imagens/escaneado)."}), 400
    titulo = request.form.get("titulo", "").strip() or re.sub(r"\.[^.]+$", "", nome)
    citacao = request.form.get("citacao", "").strip() or f"{titulo}."
    categoria = request.form.get("categoria", "ICE").strip() or "ICE"
    did = db.documento_criar(titulo, citacao, categoria, nome, len(conteudo), trechos)
    return jsonify({"ok": True, "id": did, "n_trechos": len(trechos)})


@app.route("/api/documentos/<int:id_>", methods=["PATCH"])
def api_documentos_editar(id_):
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    body = request.get_json(force=True) or {}
    campos = {k: body[k] for k in ("titulo", "citacao", "categoria") if k in body}
    if "ativo" in body:
        campos["ativo"] = 1 if body["ativo"] else 0
    db.documento_atualizar(id_, **campos)
    return jsonify({"ok": True})


@app.route("/api/documentos/<int:id_>", methods=["DELETE"])
def api_documentos_excluir(id_):
    if not _e_supervisao():
        return jsonify({"erro": "sem permissão"}), 403
    db.documento_excluir(id_)
    return jsonify({"ok": True})


@app.route("/api/documentos/buscar")
def api_documentos_buscar():
    """Pré-visualização: quais trechos seriam usados para um tema."""
    q = request.args.get("q", "")
    ids = [int(x) for x in request.args.get("ids", "").split(",") if x.isdigit()]
    trechos = referencias.buscar_trechos(q, ids or None, limite=5)
    return jsonify([{"titulo": t["titulo"], "texto": t["texto"][:400]} for t in trechos])


@app.route("/api/sessao/professor", methods=["POST"])
def api_sessao_professor():
    nome = (request.get_json(silent=True) or {}).get("professor", "").strip()
    session["professor"] = nome
    return jsonify({"ok": True})


@app.route("/api/config", methods=["GET", "POST"])
def api_config():
    if request.method == "POST" and not _e_supervisao():
        return jsonify({"erro": "Apenas a supervisão pode alterar as configurações."}), 403
    if request.method == "GET":
        cfg = carregar_config()
        return jsonify({
            "provider": cfg["provider"], "model": cfg["model"], "base_url": cfg["base_url"],
            "api_key_mascarada": (cfg["api_key"][:4] + "•••" + cfg["api_key"][-4:]) if cfg["api_key"] else "",
            "configurada": bool(cfg["api_key"]),
        })
    novo = request.get_json(force=True) or {}
    atual = {}
    if CONFIG_FILE.exists():
        try:
            atual = json.loads(CONFIG_FILE.read_text("utf-8"))
        except Exception:
            atual = {}
    for k in ("provider", "model", "base_url"):
        if k in novo:
            atual[k] = novo[k].strip()
    if novo.get("api_key", "").strip():
        atual["api_key"] = novo["api_key"].strip()
    if novo.get("limpar_chave"):
        atual["api_key"] = ""
    CONFIG_FILE.write_text(json.dumps(atual, ensure_ascii=False, indent=2), "utf-8")
    return jsonify({"ok": True, "configurada": bool(atual.get("api_key"))})


if __name__ == "__main__":
    porta = int(os.getenv("PORT", "5000"))
    print(f"Assistente de Plano de Aula rodando em http://localhost:{porta}")
    app.run(host="0.0.0.0", port=porta, debug=False)
